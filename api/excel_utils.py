from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter
from django.conf import settings
from django.http import HttpResponse


def generar_excel(nombre_archivo, encabezados, filas):
    wb = openpyxl.Workbook()
    ws = wb.active

    ruta_logo = Path(settings.BASE_DIR) / "alamex_logo.png"
    fila_inicio = 1
    if ruta_logo.exists():
        img = ExcelImage(str(ruta_logo))
        img.width = 55
        img.height = 55
        ws.add_image(img, "A1")
        ws.row_dimensions[1].height = 42
        fila_inicio = 3

    for i, encabezado in enumerate(encabezados, start=1):
        ws.cell(row=fila_inicio, column=i, value=encabezado)

    for r, fila in enumerate(filas, start=fila_inicio + 1):
        for c, valor in enumerate(fila, start=1):
            ws.cell(row=r, column=c, value=valor)

    for i, _ in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 22

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    wb.save(response)
    return response
