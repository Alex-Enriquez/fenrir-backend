from rest_framework import viewsets
from equipos.models import CategoriaEquipo, Equipo, Mantenimiento
from empleados.models import Departamento, Empleado
from licencias.models import TipoLicencia, Licencia
from asignaciones.models import Asignacion
from contrasenas.models import CategoriaContrasena, RegistroContrasena
from django.contrib.auth.models import User
from proveedores.models import Proveedor
from servicios.models import ServicioContratado, PagoServicio
from infraestructura.models import Rack, DispositivoRed, PuertoRed
import calendar
from datetime import date
import secrets
import string
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAdminUser
from .excel_utils import generar_excel
from openpyxl import load_workbook
from .serializers import (
    CategoriaEquipoSerializer, EquipoSerializer, MantenimientoSerializer,
    DepartamentoSerializer, EmpleadoSerializer,
    TipoLicenciaSerializer, LicenciaSerializer,
    AsignacionSerializer,
    CategoriaContrasenaSerializer, RegistroContrasenaSerializer,
    UsuarioSistemaSerializer, ProveedorSerializer,
    ServicioContratadoSerializer, PagoServicioSerializer,
    RackSerializer, DispositivoRedSerializer, PuertoRedSerializer,
)


from .catalogos import CatalogoViewSetMixin


class CategoriaEquipoViewSet(CatalogoViewSetMixin, viewsets.ModelViewSet):
    queryset = CategoriaEquipo.objects.all()
    serializer_class = CategoriaEquipoSerializer
    relacion_uso = "equipos"
    etiqueta_uso = ("equipo", "equipos")


class EquipoViewSet(viewsets.ModelViewSet):
    queryset = Equipo.objects.all()
    serializer_class = EquipoSerializer

    @action(detail=False, methods=["get"])
    def exportar(self, request):
        equipos = self.get_queryset().select_related("categoria")
        encabezados = [
            "Categoría", "Marca", "Modelo", "N° Serie", "N° Inventario",
            "Procesador", "RAM", "Disco tipo", "Disco capacidad",
            "Estado", "Ubicación", "IMEI",
        ]
        filas = [
            [
                eq.categoria.nombre if eq.categoria else "",
                eq.marca, eq.modelo, eq.numero_serie or "", eq.numero_inventario or "",
                eq.procesador, eq.memoria_ram, eq.disco_tipo, eq.disco_capacidad,
                eq.estado, eq.ubicacion, eq.imei or "",
            ]
            for eq in equipos
        ]
        return generar_excel("equipos.xlsx", encabezados, filas)

    @action(detail=False, methods=["get"])
    def plantilla(self, request):
        encabezados = [
            "Categoria", "Marca", "Modelo", "Numero Serie", "Numero Inventario",
            "Procesador", "RAM", "Disco Tipo", "Disco Capacidad", "Estado", "Ubicacion", "IMEI",
        ]
        fila_ejemplo = [
            "Laptop", "Dell", "Latitude 5420", "ABC123", "ALX-0001",
            "Intel Core i5-1135G7", "8GB DDR4 SO-DIMM", "SSD", "256 GB",
            "disponible", "Oficina TI", "",
        ]
        return generar_excel("plantilla_equipos.xlsx", encabezados, [fila_ejemplo])

    @action(detail=False, methods=["post"])
    def importar(self, request):
        archivo = request.FILES.get("archivo")
        if not archivo:
            return Response({"error": "No se recibió ningún archivo."}, status=400)

        try:
            wb = load_workbook(archivo)
            ws = wb.active
        except Exception:
            return Response({"error": "No se pudo leer el archivo. Asegúrate de que sea un .xlsx válido."}, status=400)

        ESTADOS_VALIDOS = ["disponible", "asignado", "reparacion", "baja"]
        creados = 0
        errores = []

        from .utils_importacion import fila_de_encabezados, texto_celda
        primera = fila_de_encabezados(ws) + 1   # los datos empiezan debajo de los encabezados (con logo, en la fila 4)
        filas = list(ws.iter_rows(min_row=primera, values_only=True))
        for i, fila in enumerate(filas, start=primera):
            if not fila or all(c is None for c in fila):
                continue
            try:
                datos = (list(fila) + [None] * 12)[:12]
                (categoria_nombre, marca, modelo, numero_serie, numero_inventario,
                 procesador, ram, disco_tipo, disco_capacidad, estado, ubicacion, imei) = datos

                if not categoria_nombre:
                    errores.append(f"Fila {i}: falta la categoría.")
                    continue
                if not marca or not modelo:
                    errores.append(f"Fila {i}: falta marca o modelo.")
                    continue

                categoria, _ = CategoriaEquipo.objects.get_or_create(nombre=str(categoria_nombre).strip())

                estado_final = str(estado).strip().lower() if estado else "disponible"
                if estado_final not in ESTADOS_VALIDOS:
                    estado_final = "disponible"

                Equipo.objects.create(
                    categoria=categoria,
                    marca=str(marca).strip(),
                    modelo=str(modelo).strip(),
                    numero_serie=texto_celda(numero_serie) or None,
                    numero_inventario=texto_celda(numero_inventario) or None,
                    imei=texto_celda(imei) or None,
                    procesador=str(procesador).strip() if procesador else "",
                    memoria_ram=str(ram).strip() if ram else "",
                    disco_tipo=str(disco_tipo).strip() if disco_tipo else "",
                    disco_capacidad=str(disco_capacidad).strip() if disco_capacidad else "",
                    estado=estado_final,
                    ubicacion=str(ubicacion).strip() if ubicacion else "",
                )
                creados += 1
            except Exception as e:
                errores.append(f"Fila {i}: {str(e)}")

        return Response({"creados": creados, "errores": errores})


class MantenimientoViewSet(viewsets.ModelViewSet):
    queryset = Mantenimiento.objects.all()
    serializer_class = MantenimientoSerializer

    @action(detail=False, methods=["get"])
    def exportar(self, request):
        mantenimientos = self.get_queryset().select_related("equipo")
        encabezados = ["Fecha", "Equipo", "Descripción", "Solución", "Realizado por", "Costo"]
        filas = [
            [
                m.fecha.isoformat() if m.fecha else "",
                str(m.equipo) if m.equipo else "",
                m.descripcion, m.solucion or "", m.realizado_por or "",
                str(m.costo) if m.costo else "",
            ]
            for m in mantenimientos
        ]
        return generar_excel("mantenimientos.xlsx", encabezados, filas)


class DepartamentoViewSet(CatalogoViewSetMixin, viewsets.ModelViewSet):
    queryset = Departamento.objects.all()
    serializer_class = DepartamentoSerializer
    relacion_uso = "empleados"
    etiqueta_uso = ("empleado", "empleados")


class EmpleadoViewSet(viewsets.ModelViewSet):
    queryset = Empleado.objects.all()
    serializer_class = EmpleadoSerializer

    @action(detail=False, methods=["get"])
    def exportar(self, request):
        empleados = self.get_queryset().select_related("departamento")
        encabezados = ["Nombre completo", "Correo", "Departamento", "Puesto", "Estatus", "Teléfono", "Fecha ingreso"]
        filas = [
            [
                e.nombre_completo, e.correo,
                e.departamento.nombre if e.departamento else "",
                e.puesto, e.estatus, e.telefono or "",
                e.fecha_ingreso.isoformat() if e.fecha_ingreso else "",
            ]
            for e in empleados
        ]
        return generar_excel("empleados.xlsx", encabezados, filas)


class TipoLicenciaViewSet(CatalogoViewSetMixin, viewsets.ModelViewSet):
    queryset = TipoLicencia.objects.all()
    serializer_class = TipoLicenciaSerializer
    relacion_uso = "licencias"
    etiqueta_uso = ("licencia", "licencias")


class LicenciaViewSet(viewsets.ModelViewSet):
    queryset = Licencia.objects.all()
    serializer_class = LicenciaSerializer

    @action(detail=False, methods=["get"])
    def exportar(self, request):
        licencias = self.get_queryset().select_related("tipo")
        encabezados = ["Tipo", "Etiqueta", "Modalidad", "Asientos", "Vencimiento", "Proveedor", "Costo"]
        filas = [
            [
                l.tipo.nombre if l.tipo else "",
                l.etiqueta or "", l.modalidad, l.asientos_totales,
                l.fecha_vencimiento.isoformat() if l.fecha_vencimiento else "",
                l.proveedor_compra or "", str(l.costo) if l.costo else "",
            ]
            for l in licencias
        ]
        return generar_excel("licencias.xlsx", encabezados, filas)


class AsignacionViewSet(viewsets.ModelViewSet):
    queryset = Asignacion.objects.all()
    serializer_class = AsignacionSerializer

    @action(detail=False, methods=["get"])
    def exportar(self, request):
        asignaciones = self.get_queryset().select_related("empleado", "equipo", "licencia")
        encabezados = ["Empleado", "Equipo", "Licencia", "Fecha asignación", "Fecha devolución"]
        filas = [
            [
                str(a.empleado) if a.empleado else "",
                str(a.equipo) if a.equipo else "",
                str(a.licencia) if a.licencia else "",
                a.fecha_asignacion.isoformat() if a.fecha_asignacion else "",
                a.fecha_devolucion.isoformat() if a.fecha_devolucion else "",
            ]
            for a in asignaciones
        ]
        return generar_excel("asignaciones.xlsx", encabezados, filas)

class CategoriaContrasenaViewSet(CatalogoViewSetMixin, viewsets.ModelViewSet):
    queryset = CategoriaContrasena.objects.all()
    serializer_class = CategoriaContrasenaSerializer
    permission_classes = [IsAdminUser]
    relacion_uso = "registros"
    etiqueta_uso = ("contraseña", "contraseñas")

class RegistroContrasenaViewSet(viewsets.ModelViewSet):
    queryset = RegistroContrasena.objects.all()
    serializer_class = RegistroContrasenaSerializer
    permission_classes = [IsAdminUser]


def generar_password_automatica():
    alfabeto = string.ascii_letters + string.digits + "!@#$%"
    return "".join(secrets.choice(alfabeto) for _ in range(12))


class UsuarioSistemaViewSet(viewsets.ModelViewSet):
    queryset = User.objects.all().order_by("username")
    serializer_class = UsuarioSistemaSerializer
    permission_classes = [IsAdminUser]

    @action(detail=True, methods=["post"])
    def reset_password(self, request, pk=None):
        user = self.get_object()
        modo = request.data.get("modo", "auto")

        if modo == "manual":
            nueva_password = request.data.get("password", "")
            if not nueva_password:
                return Response({"error": "Debes escribir una contraseña."}, status=400)
        else:
            nueva_password = generar_password_automatica()

        user.set_password(nueva_password)
        user.save()
        return Response({"password": nueva_password})

class ProveedorViewSet(viewsets.ModelViewSet):
    queryset = Proveedor.objects.all().order_by("nombre")
    serializer_class = ProveedorSerializer


def sumar_meses(fecha, meses):
    mes = fecha.month - 1 + meses
    anio = fecha.year + mes // 12
    mes = mes % 12 + 1
    dia = min(fecha.day, calendar.monthrange(anio, mes)[1])
    return date(anio, mes, dia)


MESES_POR_PERIODICIDAD = {
    "mensual": 1,
    "trimestral": 3,
    "semestral": 6,
    "anual": 12,
}


class ServicioContratadoViewSet(viewsets.ModelViewSet):
    queryset = ServicioContratado.objects.all().order_by("nombre")
    serializer_class = ServicioContratadoSerializer

    @action(detail=True, methods=["get"])
    def pagos(self, request, pk=None):
        servicio = self.get_object()
        pagos = servicio.pagos.all().order_by("-fecha")
        return Response(PagoServicioSerializer(pagos, many=True).data)

    @action(detail=True, methods=["post"])
    def marcar_pago(self, request, pk=None):
        servicio = self.get_object()
        monto = request.data.get("monto", servicio.costo)

        PagoServicio.objects.create(servicio=servicio, monto=monto)

        base = servicio.fecha_proximo_pago or date.today()
        meses = MESES_POR_PERIODICIDAD.get(servicio.periodicidad, 1)
        servicio.fecha_proximo_pago = sumar_meses(base, meses)
        servicio.save()

        return Response(ServicioContratadoSerializer(servicio).data)


class RackViewSet(viewsets.ModelViewSet):
    queryset = Rack.objects.all().order_by("nombre")
    serializer_class = RackSerializer


class DispositivoRedViewSet(viewsets.ModelViewSet):
    queryset = DispositivoRed.objects.all().order_by("nombre")
    serializer_class = DispositivoRedSerializer

    def perform_create(self, serializer):
        dispositivo = serializer.save()
        if dispositivo.tipo == "switch" and dispositivo.numero_puertos > 0:
            puertos = [
                PuertoRed(dispositivo=dispositivo, numero=n)
                for n in range(1, dispositivo.numero_puertos + 1)
            ]
            PuertoRed.objects.bulk_create(puertos)

    @action(detail=True, methods=["get"])
    def puertos(self, request, pk=None):
        dispositivo = self.get_object()
        return Response(PuertoRedSerializer(dispositivo.puertos.all(), many=True).data)


class PuertoRedViewSet(viewsets.ModelViewSet):
    queryset = PuertoRed.objects.all()
    serializer_class = PuertoRedSerializer
